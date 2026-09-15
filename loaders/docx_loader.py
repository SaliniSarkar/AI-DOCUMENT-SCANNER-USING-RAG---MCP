from docx import Document


def load_docx(uploaded_file):
    document = Document(uploaded_file)
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"Table {table_index}:")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)
